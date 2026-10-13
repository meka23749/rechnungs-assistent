"""Excel-Export: Blatt 'Rechnungen' für die Buchhaltung, Blatt 'Prüfprotokoll' zur Nachvollziehbarkeit."""

from __future__ import annotations

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EURO = '#,##0.00 "€"'
STATUS_TEXT = {"ok": "In Ordnung", "pruefen": "Prüfen", "fehler": "Fehler"}
STATUS_FARBE = {"ok": "D8EFDF", "pruefen": "FBEBC8", "fehler": "F6D5D2"}
KOPF = PatternFill("solid", fgColor="1B2A41")


def _kopfzeile(blatt, titel: list[str], breiten: list[int]) -> None:
    blatt.append(titel)
    for spalte, breite in enumerate(breiten, start=1):
        zelle = blatt.cell(row=1, column=spalte)
        zelle.font = Font(bold=True, color="FFFFFF")
        zelle.fill = KOPF
        zelle.alignment = Alignment(vertical="center")
        blatt.column_dimensions[get_column_letter(spalte)].width = breite
    blatt.freeze_panes = "A2"


def erstelle_excel(rechnungen: list[dict]) -> bytes:
    mappe = Workbook()

    blatt = mappe.active
    blatt.title = "Rechnungen"
    _kopfzeile(
        blatt,
        ["Lieferant", "Rechnungsnr.", "Rechnungsdatum", "Netto", "MwSt %", "MwSt", "Brutto",
         "IBAN", "Status", "Freigegeben", "Quelle", "Datei"],
        [32, 16, 17, 14, 9, 13, 14, 28, 13, 12, 10, 34],
    )
    # Druck: Querformat, alle Spalten auf einer Seitenbreite
    blatt.page_setup.orientation = "landscape"
    blatt.sheet_properties.pageSetUpPr.fitToPage = True
    blatt.page_setup.fitToWidth = 1
    blatt.page_setup.fitToHeight = 0
    for r in rechnungen:
        datum = date.fromisoformat(r["rechnungsdatum"]) if r.get("rechnungsdatum") else None
        blatt.append([
            r.get("lieferant"), r.get("rechnungsnummer"), datum, r.get("netto"),
            (r["mwst_satz"] / 100) if r.get("mwst_satz") is not None else None,
            r.get("mwst_betrag"), r.get("brutto"), r.get("iban"),
            STATUS_TEXT.get(r["status"], r["status"]), "Ja" if r["freigegeben"] else "Nein",
            "KI" if r["quelle"] == "ki" else "Regeln", r["dateiname"],
        ])
        zeile = blatt.max_row
        blatt.cell(zeile, 3).number_format = "DD.MM.YYYY"
        blatt.cell(zeile, 5).number_format = "0%"
        for spalte in (4, 6, 7):
            blatt.cell(zeile, spalte).number_format = EURO
        blatt.cell(zeile, 9).fill = PatternFill("solid", fgColor=STATUS_FARBE.get(r["status"], "FFFFFF"))

    if rechnungen:
        letzte = blatt.max_row
        blatt.append(["Summe", None, None, f"=SUM(D2:D{letzte})", None,
                      f"=SUM(F2:F{letzte})", f"=SUM(G2:G{letzte})"])
        summe = blatt.max_row
        for spalte in (1, 4, 6, 7):
            blatt.cell(summe, spalte).font = Font(bold=True)
        for spalte in (4, 6, 7):
            blatt.cell(summe, spalte).number_format = EURO
        blatt.auto_filter.ref = f"A1:L{letzte}"

    protokoll = mappe.create_sheet("Prüfprotokoll")
    _kopfzeile(protokoll, ["Datei", "Rechnungsnr.", "Regel", "Ergebnis", "Meldung"], [34, 16, 16, 10, 90])
    farben = {"ok": "D8EFDF", "warnung": "FBEBC8", "fehler": "F6D5D2"}
    for r in rechnungen:
        for p in r["pruefung"]:
            protokoll.append([r["dateiname"], r.get("rechnungsnummer"), p["regel"], p["status"], p["meldung"]])
            protokoll.cell(protokoll.max_row, 4).fill = PatternFill("solid", fgColor=farben.get(p["status"], "FFFFFF"))

    info = mappe.create_sheet("Info")
    info.append(["Erstellt am", datetime.now().strftime("%d.%m.%Y %H:%M")])
    info.append(["Erstellt mit", "Rechnungs-Assistent (Demo-Projekt)"])
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 40

    puffer = io.BytesIO()
    mappe.save(puffer)
    return puffer.getvalue()
